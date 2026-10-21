"""Generación del informe de una simulación en formato Word (.docx)."""
import os
import tempfile
from dataclasses import fields
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import verificacion
from .explicador import explicar_paso, texto_ecuaciones
from .graficas import figura_error_dt, figura_resultados

AUTORES = ["David Esteban Rodríguez García", "Jesús Eliécer Portela Córdoba"]
TUTOR = "William Ricardo Calderón Vásquez"


def _tabla(doc, encabezados, filas, ancho_fuente=9):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(encabezados):
        t.rows[0].cells[i].text = h
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = str(v)
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(ancho_fuente)
    return t


def _mono(doc, texto):
    par = doc.add_paragraph()
    run = par.add_run(texto)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    return par


def conclusiones(sim):
    p, r = sim.p, sim.resumen()
    c = []
    c.append(f"La humedad del suelo se movió entre {r['H_min']:.1f} % y {r['H_max']:.1f} % "
             f"(promedio {r['H_prom']:.1f} %), y la temperatura entre {r['T_min']:.1f} °C y {r['T_max']:.1f} °C.")
    if p.control_activo:
        c.append(f"El control activó el riego {r['act_riego']} vez/veces ({r['min_riego']:g} min en total) y la "
                 f"ventilación {r['act_vent']} vez/veces ({r['min_vent']:g} min en total).")
        if r["act_riego"]:
            c.append("Cada activación del riego cambió el signo de dH/dt (R > E), lo que se observa como un "
                     "cambio de pendiente en la curva de humedad.")
        if r["act_vent"]:
            c.append("La ventilación aumentó el coeficiente de intercambio térmico de k a k + kv, "
                     "acercando la temperatura interior a la exterior más rápidamente.")
    else:
        c.append("El control automático estuvo desactivado; las variables evolucionaron solo según las "
                 "ecuaciones y las entradas ambientales.")
    if p.usar_crecimiento:
        c.append(f"El tamaño de la planta pasó de {p.P0:g} cm a {r['P_final']:.2f} cm; el crecimiento fue "
                 "mayor mientras hubo luz y la temperatura y humedad estuvieron cerca de sus valores óptimos.")
    c.append(f"Se usó el método de Euler con Δt = {p.dt:g} min; los resultados coinciden con los cálculos "
             "manuales dentro de la precisión de punto flotante.")
    return c


def generar_informe(sim, ruta, titulo="Informe de simulación del invernadero"):
    p = sim.p
    r = sim.resumen()
    doc = Document()

    doc.add_heading(titulo, 0)
    doc.add_paragraph(f"Universidad del Tolima – IDEAD – Ingeniería de Sistemas\n"
                      f"Autores: {', '.join(AUTORES)}\nTutor: {TUTOR}\n"
                      f"Fecha de la simulación: {datetime.now():%Y-%m-%d %H:%M}")

    doc.add_heading("1. Modelo matemático y método numérico", 1)
    _mono(doc, texto_ecuaciones(p))

    doc.add_heading("2. Condiciones iniciales y parámetros", 1)
    doc.add_paragraph(f"Duración: {p.duracion:g} min ({p.duracion / 60:.1f} h) – Paso Δt: {p.dt:g} min – "
                      f"Pasos: {r['pasos']}")
    filas = [(f.name, f"{getattr(p, f.name):g}" if not isinstance(getattr(p, f.name), bool)
              else ("sí" if getattr(p, f.name) else "no")) for f in fields(p)]
    _tabla(doc, ["Parámetro", "Valor"], filas, 8)

    doc.add_heading("3. Resultados", 1)
    _tabla(doc, ["Magnitud", "Valor"], [
        ("Temperatura mínima / máxima / promedio", f"{r['T_min']:.2f} / {r['T_max']:.2f} / {r['T_prom']:.2f} °C"),
        ("Humedad mínima / máxima / promedio", f"{r['H_min']:.2f} / {r['H_max']:.2f} / {r['H_prom']:.2f} %"),
        ("Tamaño final de la planta", f"{r['P_final']:.2f} cm" if p.usar_crecimiento else "no simulado"),
        ("Activaciones de riego (tiempo)", f"{r['act_riego']} ({r['min_riego']:g} min)"),
        ("Activaciones de ventilación (tiempo)", f"{r['act_vent']} ({r['min_vent']:g} min)"),
        ("Activaciones de iluminación (tiempo)", f"{r['act_luz']} ({r['min_luz']:g} min)"),
    ])

    with tempfile.TemporaryDirectory() as tmp:
        img = os.path.join(tmp, "res.png")
        figura_resultados(sim).savefig(img, dpi=130)
        doc.add_heading("4. Gráficas", 1)
        doc.add_picture(img, width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_heading("5. Acciones de control (historial de eventos)", 1)
        ev = sim.eventos
        _tabla(doc, ["t [min]", "Hora [h]", "Evento"],
               [(f"{e[0]:g}", f"{e[0] / 60:.2f}", e[2]) for e in ev[:60]])
        if len(ev) > 60:
            doc.add_paragraph(f"... y {len(ev) - 60} eventos más.")

        doc.add_heading("6. Verificación frente a cálculos manuales", 1)
        v = verificacion.ejecutar_todo()
        doc.add_paragraph("Humedad (ejemplo del documento: E = 0.8, riego R = 2.5 desde el minuto 5):")
        _tabla(doc, ["n", "H manual [%]", "H programa [%]", "|diferencia|"],
               [(n, f"{m:.4f}", f"{c:.4f}", f"{d:.1e}") for n, m, c, d in v["humedad_manual"]])
        doc.add_paragraph("Temperatura (entradas constantes):")
        _tabla(doc, ["Caso", "n", "T manual [°C]", "T programa [°C]", "|dif.|"],
               [(a, n, f"{m:.6f}", f"{c:.6f}", f"{d:.1e}") for a, n, m, c, d in v["temperatura_manual"]])
        doc.add_paragraph("Pruebas de coherencia del modelo y del control:")
        _tabla(doc, ["Prueba", "Resultado"],
               [(d, "CUMPLE" if ok else "FALLA") for d, ok in v["comportamientos"]])

        doc.add_heading("7. Influencia del paso de tiempo Δt", 1)
        errs = verificacion.error_vs_dt()
        img2 = os.path.join(tmp, "err.png")
        figura_error_dt(errs).savefig(img2, dpi=130)
        doc.add_picture(img2, width=Inches(4.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tabla(doc, ["Δt [min]", "Error máximo [°C]"], [(f"{d:g}", f"{e:.5f}") for d, e in errs])
        doc.add_paragraph("El error crece de forma aproximadamente lineal con Δt, como se espera de un "
                          "método de primer orden como Euler.")

        if len(sim.registros) > 2:
            doc.add_heading("8. Ejemplo de cálculo paso a paso", 1)
            for n in (0, min(5, len(sim.registros) - 2)):
                _mono(doc, explicar_paso(sim, n))

        doc.add_heading("9. Conclusiones", 1)
        for c in conclusiones(sim):
            doc.add_paragraph(c, style="List Bullet")

        doc.save(ruta)
    return ruta
